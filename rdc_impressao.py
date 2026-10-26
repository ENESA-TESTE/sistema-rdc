# -*- coding: utf-8 -*-
"""
====================================================================
ENESA ENGENHARIA — MOTOR DE IMPRESSÃO E GERAÇÃO DE RDC
Gera PDF oficial A4 (fiel ao padrão da obra) e Excel preenchido
====================================================================
"""

import os
import io
import re
import datetime
import zipfile
import pandas as pd
from fpdf import FPDF

PASTA_BASE = os.path.dirname(os.path.abspath(__file__))
LOGO_PADRAO = os.path.join(PASTA_BASE, "logo_enesa.png")
if not os.path.exists(LOGO_PADRAO):
    LOGO_PADRAO = os.path.join(PASTA_BASE, "logo.png")
MODELO_XLSX_PADRAO = os.path.join(PASTA_BASE, "MODELO.xlsx")


def safe_text(text):
    """Garante compatibilidade de encoding com FPDF Latin-1."""
    if text is None:
        return ""
    return str(text).encode("latin-1", "replace").decode("latin-1")


def gerar_pdf_rdc(equipe, encarregado_selecionado, nome_empresa="ENESA Engenharia", logo_path="", data_rdc=""):
    """Gera um PDF que replica fielmente o MODELO oficial da Apropriação de Campo ENESA em folha A4."""
    try:
        pdf = FPDF(orientation="P", unit="mm", format="A4")
        pdf.add_page()
        pdf.set_auto_page_break(auto=False, margin=5)
        
        ml = 5        # margem esquerda
        page_w = 200  # largura útil
        
        # Cabeçalho
        y_start = 5.0
        logo_w = 28.0
        info_w = page_w - logo_w  # 172 mm
        h_row = 5.2
        h_titulo = 6.8
        header_h = h_titulo + (h_row * 4)  # 27.6 mm
        
        # Box Logo
        caminho_logo = logo_path if logo_path and os.path.exists(logo_path) else LOGO_PADRAO
        pdf.set_xy(ml, y_start)
        pdf.cell(logo_w, header_h, "", 1, 0)
        
        if os.path.exists(caminho_logo):
            try:
                pdf.image(caminho_logo, x=ml + 2, y=y_start + 2, w=logo_w - 4, h=header_h - 4)
            except Exception:
                pass
        
        # Título: Apropriação de Campo
        pdf.set_xy(ml + logo_w, y_start)
        pdf.set_font("Helvetica", "B", 9.5)
        pdf.cell(info_w, h_titulo, "APROPRIACAO DE CAMPO", 1, 1, "C")
        
        y = y_start + h_titulo
        x_info = ml + logo_w
        col_left = 75
        col_mid = 60
        col_right = info_w - col_left - col_mid  # 37 mm
        
        # Linha 2: Obra | Disciplina | Data
        pdf.set_xy(x_info, y)
        pdf.set_font("Helvetica", "", 5.5)
        pdf.cell(12, h_row, "OBRA:", 1, 0, "R")
        pdf.set_font("Helvetica", "", 7)
        obra_txt = safe_text(f" 125 - {nome_empresa}") if nome_empresa else " 125 - ARAUCO"
        pdf.cell(col_left - 12, h_row, obra_txt, 1, 0, "L")
        pdf.set_font("Helvetica", "", 5.5)
        pdf.cell(28, h_row, "DISCIPLINA/SETOR:", 1, 0, "R")
        pdf.cell(col_mid - 28, h_row, "", 1, 0, "L")
        pdf.cell(10, h_row, "DATA:", 1, 0, "R")
        pdf.set_font("Helvetica", "B", 7)
        pdf.cell(col_right - 10, h_row, str(data_rdc) if data_rdc else "", 1, 1, "C")
        
        y += h_row
        
        # Linha 3: Turno | Coordenador/Supervisor
        pdf.set_xy(x_info, y)
        pdf.set_font("Helvetica", "", 5.5)
        pdf.cell(12, h_row, "TURNO:", 1, 0, "R")
        pdf.cell(col_left - 12, h_row, safe_text(" 1o Turno (  )   2o Turno (  )   3o Turno (  )"), 1, 0, "L")
        pdf.cell(38, h_row, "COORDENADOR / SUPERVISOR:", 1, 0, "R")
        pdf.cell(info_w - col_left - 38, h_row, "", 1, 1, "L")
        
        y += h_row
        
        # Linha 4: Encarregado ou Mestre
        pdf.set_xy(x_info, y)
        pdf.set_font("Helvetica", "", 5.5)
        pdf.cell(38, h_row, "ENCARREGADO OU MESTRE:", 1, 0, "R")
        pdf.set_font("Helvetica", "B", 8)
        pdf.cell(info_w - 38, h_row, safe_text(f" {encarregado_selecionado}"), 1, 1, "L")
        
        y += h_row
        
        # Linha 5: Cond. Climáticas | Equipamento | Elevação
        pdf.set_xy(x_info, y)
        pdf.set_font("Helvetica", "", 5.5)
        pdf.cell(24, h_row, "COND. CLIMATICAS:", 1, 0, "L")
        pdf.cell(15, h_row, "Bom (  )", 1, 0, "C")
        pdf.cell(18, h_row, "Chuva Leve(  )", 1, 0, "C")
        pdf.cell(18, h_row, "Chuva Forte(  )", 1, 0, "C")
        pdf.cell(20, h_row, "EQUIPAMENTO:", 1, 0, "R")
        pdf.cell(col_mid - 20, h_row, "", 1, 0, "L")
        pdf.cell(16, h_row, "ELEVACAO:", 1, 0, "R")
        pdf.cell(col_right - 16, h_row, "", 1, 1, "L")
        
        y += h_row
        
        # Separador sutil
        pdf.set_xy(ml, y)
        pdf.cell(page_w, 1.5, "", 0, 1)
        y = pdf.get_y()
        
        # Cabeçalho da Tabela de Efetivo
        col_item = 8
        col_mat = 20
        col_nome = 60
        col_func = 45
        col_h_ini = 11
        col_h_ter = 11
        col_num = 6
        col_total = 9
        
        pdf.set_xy(ml, y)
        pdf.set_font("Helvetica", "B", 6)
        pdf.set_fill_color(225, 225, 225)
        
        pdf.cell(col_item, 3.8, "ITEM", "LTR", 0, "C", True)
        pdf.cell(col_mat, 3.8, "MATRICULA", "LTR", 0, "C", True)
        pdf.cell(col_nome, 3.8, "NOME", "LTR", 0, "C", True)
        pdf.cell(col_func, 3.8, "FUNCAO", "LTR", 0, "C", True)
        pdf.cell(col_h_ini + col_h_ter, 3.8, "HORARIO", 1, 0, "C", True)
        for n in range(1, 7):
            pdf.cell(col_num, 3.8, str(n), 1, 0, "C", True)
        pdf.cell(col_total, 3.8, "TOTAL", 1, 1, "C", True)
        
        y = pdf.get_y()
        pdf.set_xy(ml, y)
        pdf.set_font("Helvetica", "", 5)
        pdf.cell(col_item, 2.8, "", "LBR", 0, "C", True)
        pdf.cell(col_mat, 2.8, "", "LBR", 0, "C", True)
        pdf.cell(col_nome, 2.8, "", "LBR", 0, "C", True)
        pdf.cell(col_func, 2.8, "", "LBR", 0, "C", True)
        pdf.cell(col_h_ini, 2.8, "INICIO", 1, 0, "C", True)
        pdf.cell(col_h_ter, 2.8, "TERMINO", 1, 0, "C", True)
        for _ in range(6):
            pdf.cell(col_num, 2.8, "HN", 1, 0, "C", True)
        pdf.cell(col_total, 2.8, "HN", 1, 1, "C", True)
        
        y = pdf.get_y()
        
        # Linhas de Efetivo (28 linhas)
        pdf.set_font("Helvetica", "", 5.5)
        num_linhas_efetivo = 28
        h_row_efetivo = 3.75
        
        for idx in range(num_linhas_efetivo):
            pdf.set_xy(ml, y)
            if idx < len(equipe):
                row = equipe.iloc[idx]
                mat = safe_text(str(row.get("MATRICULA", "")))
                nome_col = safe_text(str(row.get("NOME", "")))
                funcao = safe_text(str(row.get("FUNCAO", row.get("FUNÇÃO", ""))))
            else:
                mat, nome_col, funcao = "", "", ""
            
            pdf.cell(col_item, h_row_efetivo, str(idx + 1), 1, 0, "C")
            pdf.cell(col_mat, h_row_efetivo, mat, 1, 0, "C")
            pdf.cell(col_nome, h_row_efetivo, nome_col, 1, 0, "L")
            pdf.cell(col_func, h_row_efetivo, funcao, 1, 0, "L")
            pdf.cell(col_h_ini, h_row_efetivo, "", 1, 0, "C")
            pdf.cell(col_h_ter, h_row_efetivo, "", 1, 0, "C")
            for _ in range(6):
                pdf.cell(col_num, h_row_efetivo, "", 1, 0, "C")
            pdf.cell(col_total, h_row_efetivo, "", 1, 1, "C")
            y = pdf.get_y()
        
        # Linha Total
        pdf.set_xy(ml, y)
        pdf.set_font("Helvetica", "B", 5.5)
        pdf.cell(col_item + col_mat + col_nome, h_row_efetivo, "", 1, 0)
        pdf.cell(col_func, h_row_efetivo, "TOTAL", 1, 0, "C")
        pdf.cell(col_h_ini, h_row_efetivo, "", 1, 0)
        pdf.cell(col_h_ter, h_row_efetivo, "", 1, 0)
        for _ in range(6):
            pdf.cell(col_num, h_row_efetivo, "", 1, 0)
        pdf.cell(col_total, h_row_efetivo, "", 1, 1)
        y = pdf.get_y()
        
        # Separador
        pdf.set_xy(ml, y)
        pdf.cell(page_w, 1.5, "", 0, 1)
        y = pdf.get_y()
        
        # Máquinas e Equipamentos
        pdf.set_xy(ml, y)
        pdf.set_font("Helvetica", "B", 5)
        pdf.cell(page_w, 3.4, "MAQUINAS E EQUIPAMENTOS (quando aplicavel)", 1, 1, "L")
        y = pdf.get_y()
        
        pdf.set_xy(ml, y)
        col_qt = 8
        col_placa = 18
        col_equip = 62
        col_tag = 30
        col_hm_ini = 12
        col_hm_ter = 12
        rest = page_w - col_qt - col_placa - col_equip - col_tag - col_hm_ini - col_hm_ter
        
        pdf.cell(col_qt, 3.2, "Qt.", 1, 0, "C", True)
        pdf.cell(col_placa, 3.2, "PLACA", 1, 0, "C", True)
        pdf.cell(col_equip, 3.2, "EQUIPAMENTO", 1, 0, "C", True)
        pdf.cell(col_tag, 3.2, "TAG", 1, 0, "C", True)
        pdf.cell(col_hm_ini, 3.2, "Inicio", 1, 0, "C", True)
        pdf.cell(col_hm_ter, 3.2, "Termino", 1, 0, "C", True)
        pdf.cell(rest, 3.2, "Total", 1, 1, "C", True)
        y = pdf.get_y()
        
        for mq in range(2):
            pdf.set_xy(ml, y)
            pdf.set_font("Helvetica", "", 5)
            pdf.cell(col_qt, 3.2, str(mq + 1), 1, 0, "C")
            pdf.cell(col_placa, 3.2, "", 1, 0)
            pdf.cell(col_equip, 3.2, "", 1, 0)
            pdf.cell(col_tag, 3.2, "", 1, 0)
            pdf.cell(col_hm_ini, 3.2, "", 1, 0)
            pdf.cell(col_hm_ter, 3.2, "", 1, 0)
            pdf.cell(rest, 3.2, "", 1, 1)
            y = pdf.get_y()
        
        # Separador
        pdf.set_xy(ml, y)
        pdf.cell(page_w, 1.5, "", 0, 1)
        y = pdf.get_y()
        
        # Atividades (15 linhas)
        pdf.set_xy(ml, y)
        pdf.set_font("Helvetica", "B", 6)
        col_item_a = 8
        col_ativ = page_w - col_item_a - 20
        col_ca = 20
        
        pdf.cell(col_item_a, 4.2, "ITEM", 1, 0, "C", True)
        pdf.cell(col_ativ, 4.2, "ATIVIDADES", 1, 0, "C", True)
        pdf.cell(col_ca, 4.2, "C. ATIVID.", 1, 1, "C", True)
        y = pdf.get_y()
        
        pdf.set_font("Helvetica", "", 5.5)
        h_row_ativ = 4.15
        for ai in range(15):
            pdf.set_xy(ml, y)
            pdf.cell(col_item_a, h_row_ativ, str(ai + 1), 1, 0, "C")
            pdf.cell(col_ativ, h_row_ativ, "", 1, 0)
            pdf.cell(col_ca, h_row_ativ, "", 1, 1)
            y = pdf.get_y()
        
        # Separador
        pdf.set_xy(ml, y)
        pdf.cell(page_w, 1.5, "", 0, 1)
        y = pdf.get_y()
        
        # Área / Local
        pdf.set_xy(ml, y)
        pdf.set_font("Helvetica", "B", 6)
        pdf.cell(page_w, 4.5, "          PB (  )            RB (  )                                                                                    OBS", 1, 1, "L")
        y = pdf.get_y()
        
        areas = [
            "DUTO (  )      EQUIPAMENTO (  )     TUBULACAO (  )",
            "ESTRUTURA MET (  )      PRECIPITADOR (  )",
            "PRESSAO - MEC (  )      PRESSAO - TUBULACAO (  )",
            "PRESSAO - FORNALHA (  )      PINTURA (  )",
            "SOPRAGEM (  )      ANDAIME (  )",
        ]
        pdf.set_font("Helvetica", "", 5.8)
        for area_txt in areas:
            pdf.set_xy(ml, y)
            pdf.cell(page_w, 3.8, safe_text(area_txt), 1, 1, "L")
            y = pdf.get_y()
        
        pdf.set_xy(ml, y)
        pdf.set_font("Helvetica", "B", 5)
        pdf.cell(page_w, 3.8, " MARCAR CONFORME O LOCAL DE TRABALHO", 1, 1, "L")
        y = pdf.get_y()
        
        # Assinaturas
        col_ass = page_w / 3
        y_linhas_ass = y + 10.0
        x1 = ml + 6
        x2 = ml + col_ass + 6
        x3 = ml + 2 * col_ass + 6
        line_len = col_ass - 12
        
        pdf.line(x1, y_linhas_ass, x1 + line_len, y_linhas_ass)
        pdf.line(x2, y_linhas_ass, x2 + line_len, y_linhas_ass)
        pdf.line(x3, y_linhas_ass, x3 + line_len, y_linhas_ass)
        
        pdf.set_xy(ml, y_linhas_ass + 1.2)
        pdf.set_font("Helvetica", "", 5.5)
        pdf.cell(col_ass, 4, "ENCARREGADO / MESTRE", 0, 0, "C")
        pdf.cell(col_ass, 4, "ENGENHEIRO / COORDENADOR", 0, 0, "C")
        pdf.cell(col_ass, 4, "PLANEJAMENTO", 0, 1, "C")
        
        return bytes(pdf.output())
    except Exception as e:
        print(f"Erro em gerar_pdf_rdc: {e}")
        return None


def preencher_excel(equipe, encarregado_selecionado, data_rdc="", modelo_path=""):
    """Preenche o arquivo MODELO.xlsx com o encarregado e equipe."""
    try:
        import openpyxl
        from copy import copy
        
        caminho_modelo = modelo_path if modelo_path and os.path.exists(modelo_path) else MODELO_XLSX_PADRAO
        if not os.path.exists(caminho_modelo):
            return None
            
        wb = openpyxl.load_workbook(caminho_modelo)
        ws = wb.active
        
        caminho_logo = LOGO_PADRAO
        if os.path.exists(caminho_logo):
            try:
                from openpyxl.drawing.image import Image as OpenpyxlImage
                _logo = OpenpyxlImage(caminho_logo)
                _logo.width = 82
                _logo.height = 108
                ws.add_image(_logo, "A1")
            except Exception:
                pass
                
        celula_enc = ws["I4"]
        celula_enc.value = encarregado_selecionado
        
        tamanho = 16
        if len(encarregado_selecionado) > 25:
            tamanho = 14
        if len(encarregado_selecionado) > 35:
            tamanho = 12
            
        if celula_enc.font:
            nova_fonte = copy(celula_enc.font)
            nova_fonte.size = tamanho
            nova_fonte.bold = True
            celula_enc.font = nova_fonte
            
        linha_atual = 9
        for _, row in equipe.iterrows():
            ws[f"B{linha_atual}"] = str(row.get("MATRICULA", ""))
            ws[f"C{linha_atual}"] = str(row.get("NOME", ""))
            ws[f"H{linha_atual}"] = str(row.get("FUNCAO", row.get("FUNÇÃO", "")))
            linha_atual += 1
            
        ws.print_area = "A1:R61"
        ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.horizontalCentered = True
        ws.page_setup.verticalCentered = True
        ws.page_margins.left = 0.2
        ws.page_margins.right = 0.2
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.5
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        
        return wb
    except Exception as e:
        print(f"Erro em preencher_excel: {e}")
        return None
